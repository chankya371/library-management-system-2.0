import csv
import datetime
import os
from typing import Dict, Any, List, Tuple
from database.db_connection import get_db_cursor
from services.circulation_service import CirculationService
from services.fine_service import FineService

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class ReportService:
    """Computes analytical reports, chart data, and file exports."""

    @staticmethod
    def get_dashboard_summary() -> Dict[str, Any]:
        """Provides real-time numbers for all dashboard cards."""
        CirculationService.sync_overdue_statuses()

        with get_db_cursor(commit=False) as cursor:
            # Books stats
            cursor.execute("SELECT COALESCE(SUM(total_copies), 0), COALESCE(SUM(available_copies), 0) FROM books;")
            total_copies, available_copies = cursor.fetchone()
            issued_copies = total_copies - available_copies

            # Distinct book titles count
            cursor.execute("SELECT COUNT(*) FROM books;")
            total_titles = cursor.fetchone()[0]

            # Members
            cursor.execute("SELECT COUNT(*) FROM members WHERE status = 'Active';")
            active_members = cursor.fetchone()[0]
            cursor.execute("SELECT COUNT(*) FROM members;")
            total_members = cursor.fetchone()[0]

            # Overdue issues
            cursor.execute("SELECT COUNT(*) FROM issue_records WHERE status = 'Overdue';")
            overdue_count = cursor.fetchone()[0]

        collected_fines, pending_fines = FineService.get_fine_totals()

        return {
            "total_books": total_copies,
            "total_titles": total_titles,
            "available_books": available_copies,
            "issued_books": issued_copies,
            "total_members": total_members,
            "active_members": active_members,
            "overdue_books": overdue_count,
            "total_fines_collected": collected_fines,
            "total_fines_pending": pending_fines,
        }

    @staticmethod
    def get_circulation_trend_data(days: int = 7) -> Dict[str, Any]:
        """Returns issued vs returned count for the last N days."""
        today = datetime.date.today()
        dates = [(today - datetime.timedelta(days=i)).isoformat() for i in range(days - 1, -1, -1)]
        day_labels = [(today - datetime.timedelta(days=i)).strftime("%a %d") for i in range(days - 1, -1, -1)]

        issued_counts = []
        returned_counts = []

        with get_db_cursor(commit=False) as cursor:
            for dt in dates:
                cursor.execute("SELECT COUNT(*) FROM issue_records WHERE issue_date = ?;", (dt,))
                issued_counts.append(cursor.fetchone()[0])

                cursor.execute(
                    "SELECT COUNT(*) FROM issue_records WHERE return_date LIKE ?;",
                    (f"{dt}%",),
                )
                returned_counts.append(cursor.fetchone()[0])

        return {
            "labels": day_labels,
            "issued": issued_counts,
            "returned": returned_counts,
        }

    @staticmethod
    def get_category_distribution() -> Dict[str, int]:
        """Returns count of books grouped by category."""
        with get_db_cursor(commit=False) as cursor:
            cursor.execute(
                """
                SELECT category, SUM(total_copies) as total 
                FROM books 
                GROUP BY category 
                ORDER BY total DESC;
                """
            )
            rows = cursor.fetchall()
            return {r["category"]: r["total"] for r in rows}

    @staticmethod
    def get_most_borrowed_books(limit: int = 5) -> List[Dict[str, Any]]:
        with get_db_cursor(commit=False) as cursor:
            cursor.execute(
                """
                SELECT b.id, b.title, b.author, b.category, COUNT(ir.id) as borrow_count
                FROM books b
                LEFT JOIN issue_records ir ON b.id = ir.book_id
                GROUP BY b.id
                ORDER BY borrow_count DESC, b.title ASC
                LIMIT ?;
                """,
                (limit,),
            )
            return [dict(r) for r in cursor.fetchall()]

    @staticmethod
    def get_most_active_members(limit: int = 5) -> List[Dict[str, Any]]:
        with get_db_cursor(commit=False) as cursor:
            cursor.execute(
                """
                SELECT m.id, m.full_name, m.member_code, m.membership_type, COUNT(ir.id) as total_borrows
                FROM members m
                LEFT JOIN issue_records ir ON m.id = ir.member_id
                GROUP BY m.id
                ORDER BY total_borrows DESC, m.full_name ASC
                LIMIT ?;
                """,
                (limit,),
            )
            return [dict(r) for r in cursor.fetchall()]

    @staticmethod
    def get_overdue_report() -> List[Dict[str, Any]]:
        CirculationService.sync_overdue_statuses()
        with get_db_cursor(commit=False) as cursor:
            cursor.execute(
                """
                SELECT ir.id as issue_id, b.title, b.isbn, m.full_name, m.email, m.phone, 
                       ir.issue_date, ir.due_date, ir.fine_amount
                FROM issue_records ir
                JOIN books b ON ir.book_id = b.id
                JOIN members m ON ir.member_id = m.id
                WHERE ir.status = 'Overdue'
                ORDER BY ir.due_date ASC;
                """
            )
            return [dict(r) for r in cursor.fetchall()]

    @staticmethod
    def get_recently_added_books(limit: int = 5) -> List[Dict[str, Any]]:
        with get_db_cursor(commit=False) as cursor:
            cursor.execute(
                """
                SELECT id, isbn, title, author, category, available_copies, total_copies, created_at
                FROM books 
                ORDER BY id DESC LIMIT ?;
                """,
                (limit,),
            )
            return [dict(r) for r in cursor.fetchall()]

    @staticmethod
    def get_recently_issued_books(limit: int = 5) -> List[Dict[str, Any]]:
        with get_db_cursor(commit=False) as cursor:
            cursor.execute(
                """
                SELECT ir.id, b.title, m.full_name, ir.issue_date, ir.due_date, ir.status
                FROM issue_records ir
                JOIN books b ON ir.book_id = b.id
                JOIN members m ON ir.member_id = m.id
                ORDER BY ir.id DESC LIMIT ?;
                """,
                (limit,),
            )
            return [dict(r) for r in cursor.fetchall()]

    @staticmethod
    def export_to_csv(filepath: str, headers: List[str], rows: List[List[Any]]) -> Tuple[bool, str]:
        try:
            with open(filepath, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                writer.writerows(rows)
            return True, f"Exported successfully to {os.path.basename(filepath)}"
        except Exception as e:
            return False, f"CSV Export error: {str(e)}"

    @staticmethod
    def export_to_excel(filepath: str, sheet_name: str, headers: List[str], rows: List[List[Any]]) -> Tuple[bool, str]:
        if not HAS_OPENPYXL:
            return False, "openpyxl is not installed. Export to CSV instead."

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name

            # Header styling
            header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

            ws.append(headers)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Data rows
            for row_idx, r_data in enumerate(rows, start=2):
                ws.append(r_data)
                # Alternating row background
                if row_idx % 2 == 0:
                    zebra_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
                    for col_num in range(1, len(r_data) + 1):
                        ws.cell(row=row_idx, column=col_num).fill = zebra_fill

            # Auto-fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(filepath)
            return True, f"Excel workbook exported successfully to {os.path.basename(filepath)}"
        except Exception as e:
            return False, f"Excel Export error: {str(e)}"
